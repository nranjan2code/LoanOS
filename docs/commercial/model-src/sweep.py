#!/usr/bin/env python3
"""Substitute for recalc.py: evaluate EVERY formula cell and report errors."""
import sys, os, logging
logging.disable(logging.WARNING)
from pycel import ExcelCompiler
from openpyxl import load_workbook

XLSX = sys.argv[1]
exc = ExcelCompiler(XLSX)
wb = load_workbook(XLSX)
ERRS = ("#REF!", "#VALUE!", "#NAME?", "#DIV/0!", "#N/A", "#NULL!", "#NUM!")
total = errors = blank = 0
bad = []
for s in wb.worksheets:
    for row in s.iter_rows():
        for c in row:
            if isinstance(c.value, str) and c.value.startswith("="):
                total += 1
                addr = f"'{s.title}'!{c.coordinate}"
                try:
                    v = exc.evaluate(addr)
                except Exception as e:
                    errors += 1
                    bad.append((addr, f"EXC {type(e).__name__}: {str(e)[:70]}"))
                    continue
                if isinstance(v, str) and v in ERRS:
                    errors += 1
                    bad.append((addr, v))
                elif v is None:
                    blank += 1
                    bad.append((addr, "None"))
print(f"total_formulas: {total}")
print(f"total_errors:   {errors}")
print(f"none_results:   {blank}")
if bad:
    print("\nfirst 25 problems:")
    for a, v in bad[:25]:
        print(f"  {a:<42} {v}")
print("\nstatus:", "success" if errors == 0 and blank == 0 else "errors_found")
sys.exit(1 if (errors or blank) else 0)
