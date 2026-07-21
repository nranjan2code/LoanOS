#!/usr/bin/env python3
"""Evaluate the workbook's real formulas with pycel and reconcile against the shadow model."""
import sys, os, logging
logging.disable(logging.WARNING)
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from pycel import ExcelCompiler
from openpyxl import load_workbook
from shadow import origination, servicing, recurring, admission, pnl
from build_model_part2 import JOURNEYS

XLSX = sys.argv[1]
exc = ExcelCompiler(XLSX)
wb = load_workbook(XLSX)
fails = []


def label_row(sheet, label, col="A"):
    ws = wb[sheet]
    for r in range(1, ws.max_row + 1):
        if str(ws[f"{col}{r}"].value).strip() == label.strip():
            return r
    raise KeyError(f"{sheet}!{label}")


def ev(addr):
    return exc.evaluate(addr)


def cmp(name, got, want, tol=1e-4):
    try:
        ok = abs(float(got) - float(want)) <= tol * max(1.0, abs(float(want)))
    except (TypeError, ValueError):
        ok = str(got) == str(want)
    g = round(float(got), 2) if isinstance(got, (int, float)) else got
    w = round(float(want), 2) if isinstance(want, (int, float)) else want
    print(("  ok   " if ok else "  FAIL ") + f"{name:<58} workbook={g:>17}  shadow={w:>17}")
    if not ok:
        fails.append(name)


print("=" * 120)
print("FORMULA EVALUATION — workbook formulas computed by pycel vs the independent shadow model")
print("=" * 120)

print("\n-- 7_Cost_Origination (personal_loan) --")
o = origination("personal_loan")
cmp("funded rate", ev("'7_Cost_Origination'!B9"), o["funded_rate"])
O = "7_Cost_Origination"
cmp("origination ₹ per application", ev(f"'{O}'!J{label_row(O,'TOTAL ORIGINATION EXTERNAL COST','B')}"), o["per_app"])
cmp("origination ₹ per funded loan", ev(f"'{O}'!K{label_row(O,'TOTAL ORIGINATION EXTERNAL COST','B')}"), o["per_loan"])
cmp("  platform-borne ₹/app", ev(f"'{O}'!J{label_row(O,'  of which platform-borne','B')}"), o["plat_per_app"])
cmp("  pass-through ₹/app", ev(f"'{O}'!J{label_row(O,'  of which pass-through','B')}"), o["pass_per_app"])

print("\n-- 8_Cost_Servicing (personal_loan) --")
s = servicing("personal_loan")
V = "8_Cost_Servicing"
for lbl, want in [("Effective life (months)", s["eff_months"]), ("Effective life (years)", s["eff_years"]),
                  ("EMI count over life", s["emi"]), ("Bounced EMIs over life", s["bounced"]),
                  ("Delinquent loan-months (all buckets)", sum(s["delinq"]))]:
    cmp(lbl, ev(f"'{V}'!B{label_row(V, lbl)}"), want)
for lbl, want in [("Lifetime SERVICING cost per funded loan", s["servicing"]),
                  ("Lifetime COLLECTIONS cost per funded loan", s["collections"]),
                  ("TOTAL POST-ORIGINATION COST PER FUNDED LOAN", s["total"]),
                  ("  of which platform-borne", s["plat"]),
                  ("  of which pass-through", s["passthru"]),
                  ("Cost per ACTIVE LOAN-MONTH", s["per_active_month"])]:
    cmp(lbl, ev(f"'{V}'!I{label_row(V, lbl, 'B')}"), want)

print("\n-- 3_RateCard_Recurring --")
rec, R = recurring(), "3_RateCard_Recurring"
cmp("TOTAL recurring ₹ per year", ev(f"'{R}'!L{label_row(R,'TOTAL recurring ₹ per year','B')}"), rec["total"])
cmp("  of which pass-through", ev(f"'{R}'!L{label_row(R,'  of which pass-through','B')}"), rec["passthru"])
cmp("TOTAL admission ₹ (one time)", ev(f"'{R}'!L{label_row(R,'TOTAL admission ₹ (one time)','B')}"), admission())

print("\n-- 14_Journey_Economics: all 21 journeys, workbook SUMPRODUCT vs shadow --")
print(f"       {'journey':<28}{'orig ₹/loan':>13}{'svc ₹/loan':>13}{'coll ₹/loan':>13}"
      f"{'TOTAL ₹/loan':>14}{'₹/active mo':>13}{'bps':>8}   max Δ")
worst = 0.0
for i, (jt, *_x) in enumerate(JOURNEYS):
    oj, sj = origination(jt), servicing(jt)
    row = 5 + i
    g = {k: ev(f"'14_Journey_Economics'!{c}{row}") for k, c in
         (("orig_app", "H"), ("orig", "I"), ("svc", "J"), ("col", "K"), ("tot", "L"), ("pm", "M"), ("bps", "N"))}
    want_tot = oj["per_loan"] + sj["servicing"] + sj["collections"]
    d = max(abs(g["orig_app"] - oj["per_app"]), abs(g["orig"] - oj["per_loan"]),
            abs(g["svc"] - sj["servicing"]), abs(g["col"] - sj["collections"]),
            abs(g["tot"] - want_tot), abs(g["pm"] - sj["per_active_month"]))
    worst = max(worst, d)
    st = "ok  " if d < 0.01 else "FAIL"
    print(f"  {st} {jt:<28}{g['orig']:>13,.0f}{g['svc']:>13,.0f}{g['col']:>13,.0f}"
          f"{g['tot']:>14,.0f}{g['pm']:>13,.2f}{g['bps']:>8,.0f}   {d:.5f}")
    if d >= 0.01:
        fails.append(jt)
print(f"  worst deviation across 21 journeys x 6 measures: {worst:.6f}")

print("\n-- 12_Tenant_PnL --")
p = pnl()
S = "12_Tenant_PnL"
for lbl, want in [
    ("Annual applications started", p["apps"]), ("Annual funded loans", p["F"]),
    ("Annual disbursement (₹)", p["disb"]), ("Effective life (years)", p["L"]),
    ("Steady-state active loans", p["F"] * p["L"]), ("ACTIVE LOANS (selected)", p["active"]),
    ("Servicing ramp factor", p["ramp"]),
    ("Origination external — platform-borne", p["orig_plat"]),
    ("Origination external — pass-through", p["orig_pass"]),
    ("Servicing + collections — platform-borne", p["svc_plat"]),
    ("Servicing + collections — pass-through", p["svc_pass"]),
    ("Recurring subscriptions — platform-borne", p["rec"]["plat"]),
    ("Recurring subscriptions — pass-through", p["rec"]["passthru"]),
    ("Cloud — fixed", p["cloud_fix"]), ("Cloud — variable", p["cloud_var"]),
    ("Cloud — retained storage (year 1)", p["stor_y1"]),
    ("Governance + regulatory", p["gov"]), ("Support", p["support"]),
    ("Onboarding amortised", p["onb_amort"]), ("TOTAL COST", p["cost"]),
    ("Total PASS-THROUGH cost", p["pass_total"]),
    ("Platform fee", p["rev_p"]), ("Per-decisioned-application fee", p["rev_d"]),
    ("Per-funded-loan fee", p["rev_l"]), ("LMS fee on the active book", p["rev_m"]),
    ("Pass-through recovery", p["rev_pass"]), ("TOTAL REVENUE", p["rev"]),
    ("Gross profit", p["rev"] - p["cost"]),
    ("Gross margin %", (p["rev"] - p["cost"]) / p["rev"]),
    ("Software margin % (pass-through stripped at cost)", p["soft_margin"]),
    ("Cost per funded loan (₹)", p["cost"] / p["F"]),
    ("Effective cost to tenant (bps of disbursement)", p["rev"] / p["disb"] * 10000),
    ("Cost per active loan per month (₹)", s["per_active_month"]),
    ("Annual fixed cost (₹)", p["fixed_cost"]),
    ("  of which truly volume-independent (₹)", p["fixed_indep"]),
    ("Variable cost per application (₹)", p["var_per_app"]),
    ("Revenue per application (₹)", p["rev_per_app"]),
    ("Break-even annual applications", p["be_apps"]),
]:
    cmp(lbl, ev(f"'{S}'!B{label_row(S, lbl)}"), want)

print("\n-- 13_Price_Card --")
cmp("required fee per funded loan at target GM",
    ev(f"'13_Price_Card'!B{label_row('13_Price_Card','REQUIRED fee per funded loan ₹')}"), p["req_lfee"])

print("\n-- 15_Sensitivity --")
for v in [250, 2500, 25000]:
    row = 5 + [250, 500, 1000, 2500, 5000, 10000, 25000, 50000].index(v)
    q = pnl(monthly=v)
    cost = q["fixed_indep"] + q["gov"] + q["stor_y1"] + q["var_per_app"] * q["apps"]
    rev = q["fixed_rev"] + q["rev_per_app"] * q["apps"]
    cmp(f"{v:,} apps/mo — active loans", ev(f"'15_Sensitivity'!D{row}"), q["active"])
    cmp(f"{v:,} apps/mo — governance ₹/yr", ev(f"'15_Sensitivity'!E{row}"), q["gov"])
    cmp(f"{v:,} apps/mo — storage ₹/yr", ev(f"'15_Sensitivity'!F{row}"), q["stor_y1"])
    cmp(f"{v:,} apps/mo — total cost ₹/yr", ev(f"'15_Sensitivity'!G{row}"), cost)
    cmp(f"{v:,} apps/mo — gross profit ₹/yr", ev(f"'15_Sensitivity'!I{row}"), rev - cost)

print("\n-- 16_Coverage_Map --")
cov = wb["16_Coverage_Map"]
n = sum(1 for r in range(5, cov.max_row + 1) if str(cov[f"A{r}"].value or "").startswith("INT-"))
notc = sum(1 for r in range(5, cov.max_row + 1) if cov[f"D{r}"].value == "NOT COSTED")
priced = sum(1 for r in range(5, cov.max_row + 1) if cov[f"D{r}"].value == "Priced driver")
print(f"  integrations listed {n} · priced driver {priced} · absorbed elsewhere {n-priced-notc} · NOT COSTED {notc}")
cmp("all 115 integrations listed", n, 115)
cmp("nothing left uncosted", notc, 0)

print()
print("=" * 120)
print("ALL EVALUATED FORMULAS RECONCILE" if not fails else f"{len(fails)} MISMATCH(ES): {fails}")
sys.exit(1 if fails else 0)
